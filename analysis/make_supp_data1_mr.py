"""Supplementary Data 1: per-scan outputs of the probe models, the public model and the rule-based modules,
keyed by the public identifiers of each resource; scan-level BHX box counts linked by study identifier; the model-free description of the
Seg-CQ500 masks; the adjudication case list; the registry of reported values and a sheet of formula checks.

Reference labels and masks are not included (they must be obtained from the dataset creators); reference volumes derived from
the released masks are included so that agreement statistics can be recomputed.
Usage: python make_supp_data1.py [output.xlsx]
"""
import json, sys
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
W = HERE.parent
E = W / 'data' / 'evidence'
T = W / 'data' / 'tables'
OUT = W / 'out'
PMJ = json.load(open(OUT / 'public_model_cq500.json'))
dst = Path(sys.argv[1]) if len(sys.argv) > 1 else W / 'manuscript' / 'build_mr' / 'Supplementary_Data_1.xlsx'

FONT = 'Arial'
HFILL = PatternFill('solid', start_color='E8E6DF')
thin = Side(style='thin', color='999999')


def r(x, d=6):
    try:
        return None if pd.isna(x) else round(float(x), d)
    except Exception:
        return x


# ----------------------------------------------------------------------------- CQ500 (489 analyzed scans)
cq = pd.read_csv(E / 'v3_subtypes' / 'cq500_v3b_per_scan.csv')
assert len(cq) == 489 and cq.name.is_unique
sp = cq.spacing.apply(lambda s: [float(v) for v in str(s).strip('[]').split(',')])
cq_out = pd.DataFrame({
    'scan': cq.name,
    'soft_tissue_series': cq.soft_series_x,
    'voxel_spacing_x_mm': sp.apply(lambda v: r(v[0], 4)), 'voxel_spacing_y_mm': sp.apply(lambda v: r(v[1], 4)),
    'slice_interval_mm': sp.apply(lambda v: r(v[2], 3)),
    'bone_image_source': cq.bone_source,
    'qc_gate_flag': cq.qc_flag.astype(bool),
    'v2_volume_ml': cq.pred_volume_ml.apply(r), 'v2_components': cq.pred_components,
    'v2_call_at_1ml': (cq.pred_volume_ml >= 1.0).astype(int),
    'midline_status': cq.midline_status, 'csf_offset_max_abs_mm': cq.csf_offset_maxabs_mm.apply(lambda x: r(x, 3)),
    'v3b_EDH_ml': cq.EDH_ml.apply(r), 'v3b_IPH_ml': cq.IPH_ml.apply(r), 'v3b_IVH_ml': cq.IVH_ml.apply(r), 'v3b_SAH_ml': cq.SAH_ml.apply(r),
    'v3b_SDH_ml': cq.SDH_ml.apply(r), 'v3b_any_ml': cq.ANY_ml.apply(r),
})

# ----------------------------------------------------------------------------- BHSD test (96 scans)
v1 = pd.read_csv(E / 'neurocare_neuroimaging_advanced_20261021' / 'per_scan_quantitative_metrics.csv').set_index('case_id')
v2pm = pd.read_csv(E / 'neurocare_v2_20261021' / 'hemorrhage_v2' / 'v2_metrics.csv')
v2pm_b = v2pm[v2pm.dataset == 'BHSD_test'].set_index('case_id')
v2 = pd.read_csv(E / 'neurocare_v2_20261021' / 'hemorrhage_v2' / 'ablation_no_postmask.csv').set_index('case_id')
v3b = pd.read_csv(E / 'v3_subtypes' / 'v3b' / 'v3b_bhsd_test_metrics.csv').set_index('case_id')
qc = pd.read_csv(T / 'QC_features_all_188.csv').set_index('case_id')
ids = list(v1.index)
assert len(ids) == 96 and set(ids) == set(v2pm_b.index) == set(v2.index) == set(v3b.index)
assert (abs(v1.reference_volume_ml - v2pm_b.loc[ids, 'ref_ml'].values) < 1e-6).all()
bh = pd.DataFrame({'case_id': ids})
bh['qc_p01_HU'] = [qc.loc[i, 'p01'] for i in ids]
bh['qc_gate_flag'] = [bool(qc.loc[i, 'p01'] > -500) for i in ids]
bh['reference_volume_ml'] = [r(v1.loc[i, 'reference_volume_ml']) for i in ids]
bh['v1_dice'] = [r(v1.loc[i, 'dice']) for i in ids]
bh['v1_volume_ml'] = [r(v1.loc[i, 'prediction_volume_ml']) for i in ids]
bh['v2PM_dice'] = [r(v2pm_b.loc[i, 'dice']) for i in ids]
bh['v2PM_volume_ml'] = [r(v2pm_b.loc[i, 'pred_ml']) for i in ids]
bh['v2_dice'] = [r(v2.loc[i, 'dice_nopostmask']) for i in ids]
bh['v2_volume_ml'] = [r(v2.loc[i, 'pred_ml']) for i in ids]
bh['v3b_any_dice'] = [r(v3b.loc[i, 'ANY_dice']) for i in ids]
bh['v3b_any_volume_ml'] = [r(v3b.loc[i, 'ANY_pred_ml']) for i in ids]
for c in ['EDH', 'IPH', 'IVH', 'SAH', 'SDH']:
    bh[f'{c}_reference_ml'] = [r(v3b.loc[i, f'{c}_ref_ml']) for i in ids]
    bh[f'v3b_{c}_ml'] = [r(v3b.loc[i, f'{c}_pred_ml']) for i in ids]
    bh[f'v3b_{c}_dice'] = [r(v3b.loc[i, f'{c}_dice']) for i in ids]

qc_all = pd.read_csv(T / 'QC_features_all_188.csv')
qc_out = pd.DataFrame({'case_id': qc_all.case_id, 'partition': qc_all.partition, 'p01_HU': qc_all.p01, 'min_HU': qc_all['min'],
                       'fraction_in_window': qc_all.frac_window.apply(r), 'qc_gate_flag': qc_all.p01 > -500})

# ----------------------------------------------------------------------------- Seg-CQ500 (43 scans)
s4 = pd.read_csv(T / 'SegCQ500_four_method_comparison.csv')
sv2pm = v2pm[v2pm.dataset == 'Seg-CQ500'].set_index('case_id')
sv2 = pd.read_csv(E / 'neurocare_v2_20261021' / 'hemorrhage_v2' / 'segcq500_no_postmask.csv').set_index('case_id')
sids = sorted(sv2.index, key=lambda s: int(s.split('-')[-1]))
assert len(sids) == 43 and set(sids) == set(sv2pm.index) == set(s4.case_id)
seg = pd.DataFrame({'scan': sids})
seg['reference_volume_ml'] = [r(sv2.loc[i, 'ref_ml']) for i in sids]
piv = {m: g.set_index('case_id') for m, g in s4.groupby('method')}
for m, lab in [('lightunet25d', 'v1'), ('blast_model1', 'BLASTCT'), ('blast_min100mm3', 'BLASTCT_min100mm3'), ('blast_hu45_r2', 'BLASTCT_hu45_r2')]:
    seg[f'{lab}_dice'] = [r(piv[m].loc[i, 'dice']) for i in sids]
    vol = piv[m]['prediction_volume_ml'].combine_first(piv[m]['pred_volume_ml']) if 'pred_volume_ml' in piv[m] else piv[m]['prediction_volume_ml']
    seg[f'{lab}_volume_ml'] = [r(vol.loc[i]) for i in sids]
seg['v2PM_dice'] = [r(sv2pm.loc[i, 'dice']) for i in sids]
seg['v2PM_volume_ml'] = [r(sv2pm.loc[i, 'pred_ml']) for i in sids]
seg['v2_dice'] = [r(sv2.loc[i, 'dice_nopostmask']) for i in sids]
seg['v2_volume_ml'] = [r(sv2.loc[i, 'pred_ml']) for i in sids]
cqv = cq.set_index('name')
seg['v2_volume_ml_from_CQ500_DICOM'] = [r(cqv.loc[i, 'pred_volume_ml']) if i in cqv.index else None for i in sids]

# ----------------------------------------------------------------------------- MR-BHSD (43 scans x 3 raters, v1)
mr = pd.read_csv(T / 'MRBHSD_full43_compact_model_rater_metrics.csv')
mr_out = pd.DataFrame({'case_id': mr.case_id, 'rater': mr.reference_rater, 'model': 'v1', 'dice': mr.dice.apply(r),
                       'both_masks_empty': mr.both_empty.astype(bool), 'rater_volume_ml': mr.reference_volume_ml.apply(r), 'v1_volume_ml': mr.pred_volume_ml.apply(r)})

# ----------------------------------------------------------------------------- PHE-SICH-CT-IDS
ph = pd.read_csv(E / 'neurocare_v2_20261021' / 'phe_sich' / 'phe_laterality_midline.csv', dtype={'case': str})
split = json.load(open(E / 'neurocare_v2_20261021' / 'edema' / 'split.json'))
eligible = set(split['train']) | set(split['test'])
assert len(ph) == 120 and len(eligible) == 112 and eligible <= set(ph.case)
excl = ~ph.case.isin(eligible)  # eight patients whose mask affine reversed an in-plane axis; excluded without repair
assert excl.sum() == 8
cohort = ph.case.isin(eligible) & (ph.documented_side != 'bilateral')
assert cohort.sum() == 111


def blank_excluded(series):
    return series.where(~excl, None)


ph_out = pd.DataFrame({'patient': ph.case, 'geometry_excluded': excl.values, 'laterality_cohort': cohort.values,
                       'midline_status': ph.midline_status, 'reliable_slices': ph.n_reliable_slices,
                       'midline_angle_deg': ph.midline_angle_deg, 'mean_symmetry_dice': ph.mean_symmetry_dice.apply(r),
                       'side_from_expert_edema_mask': blank_excluded(ph.edema_side_auto), 'side_from_v1_hemorrhage_mask': blank_excluded(ph.blood_v1_side_auto),
                       'expert_edema_volume_ml': blank_excluded(ph.edema_ml_assumed_mm.apply(r)), 'v1_hemorrhage_volume_ml': blank_excluded(ph.blood_v1_ml_assumed_mm.apply(r)),
                       'qc_p01_HU': ph.p01})
ed = pd.read_csv(E / 'neurocare_v2_20261021' / 'edema' / 'edema_test_metrics.csv', dtype={'case': str})
ed_out = pd.DataFrame({'patient': ed.case, 'edema_dice': ed.dice.apply(r), 'predicted_edema_ml': ed.pred_ml.apply(r), 'expert_edema_ml': ed.ref_ml.apply(r)})

# ----------------------------------------------------------------------------- registry
REG = json.load(open(OUT / 'registry.json'))
reg_rows = []


def flat(prefix, o):
    if isinstance(o, dict):
        for k, v in o.items():
            flat(f'{prefix}.{k}' if prefix else str(k), v)
    elif isinstance(o, (list, tuple)):
        reg_rows.append((prefix, json.dumps(o, ensure_ascii=False)))
    else:
        reg_rows.append((prefix, ('yes' if o else 'no') if isinstance(o, bool) else o if isinstance(o, (int, float)) else str(o)))


flat('', REG)
DROP = ('seg_ct108.labels.', 'cq_series_C')  # CQ500 reader labels of single scans; not reported values
reg_rows = [(k, v) for k, v in reg_rows if not k.startswith(DROP)]
assert not any('ref_ICH' in str(v) for _, v in reg_rows)
reg_out = pd.DataFrame(reg_rows, columns=['key', 'value'])

# ----------------------------------------------------------------------------- workbook
README = [
    ('Supplementary Data 1', 'Per-scan outputs of the probe models, the public model and the rule-based modules, and the registry of reported values'),
    ('Manuscript', 'Openly downloadable head CT datasets provide insufficient reference support for validating key intracranial hemorrhage measurements: an endpoint-specific assessment'),
    ('Contents', 'One sheet per resource, keyed by the public identifiers of that resource; a Registry sheet; a Checks sheet that recomputes headline values with formulas. '
                 'Dice coefficients and volumes are stored to six decimals.'),
    ('Not included', 'Images, reference masks and reference labels (CQ500 reads, BHX boxes, PHE-SICH-CT-IDS documented side). The CQ500 reads (the creators’ reads.csv, CC BY-NC-SA 4.0) are also '
                     'provided in the code repository; obtain the other references from their creators under their licenses (Data availability) and join on the identifier columns. '
                     'Reference volumes derived from the released masks are included so that agreement statistics can be recomputed.'),
    ('CQ500', '489 analyzed CQ500 scans (scan = folder name CQ500-CT-n). Selected non-contrast soft-tissue series, voxel spacing, QC gate flag, v2 predicted volume (locked model, no post-mask) '
              'and component count, the prespecified call at 1.0 mL, midline status, the exploratory CSF-offset distance (maximum absolute offset over reliable slices; not a midline-shift '
              'measurement) and v3b class volumes (post hoc model). Join with the creators’ reads.csv (R1–R3 labels; also in the code repository) on scan.'),
    ('BHSD_test', '96 BHSD test scans: QC first-percentile stored intensity and gate flag (abstain when p01 > −500), reference volume, Dice and predicted volume for v1, v2-PM (prespecified), '
                  'v2 (post hoc, no post-mask) and v3b (post hoc), and per-class reference volume, v3b class volume and class Dice (blank when both masks are empty).'),
    ('QC_features', 'QC gate features for all 188 BHSD scans (92 training, 96 test).'),
    ('SegCQ500', '43 Seg-CQ500 scans: reference volume, Dice and volume for v1, the pretrained BLAST-CT model and two BLAST-CT post-processing variants, v2-PM and v2, and the v2 volume '
                 'computed from the CQ500 DICOM series of the same scan (cross-pipeline check).'),
    ('MRBHSD', '43 MR-BHSD validation scans × 3 raters: Dice of v1 against each rater (set to 1 when both masks are empty, as in the original analysis), rater volume and v1 volume.'),
    ('PHE_laterality', '120 PHE-SICH-CT-IDS patients: geometry_excluded marks the eight patients (0160–0167) whose mask affine reversed an in-plane axis and who were excluded without repair '
                       '(their side calls and volumes are left blank); laterality_cohort marks the 111 geometry-eligible patients with unilateral documented hemorrhage, the denominator of the '
                       'laterality analysis before abstentions. Also ideal-midline status and symmetry, side derived from the expert edema mask and from the v1 hemorrhage mask, volumes and QC '
                       'intensity. Volumes assume millimeter header units.'),
    ('PHE_edema_test', '37 held-out PHE-SICH-CT-IDS patients: edema-model Dice and volumes.'),
    ('PublicModel', 'The public model (ianpan/ct-head-hemorrhage-detection, revision e6d89c2; RSNA 2019-trained) on the 489 analyzed CQ500 scans: selected series, slices, '
                    'series-level fixed-blend probabilities (p_*), base (base_*) and contextual (ctx_*) probabilities and the maximum slice probability (slicemax_*) for any hemorrhage and '
                    'the five subtypes. Post hoc replication; plan locked before the evaluation outputs were generated.'),
    ('CQ500_UIDs', 'DICOM study and series identifiers of every series in the 489 analyzed CQ500 archives, read from the released headers; locked_soft_series marks the series analyzed.'),
    ('BHX_scan_counts', 'Number of BHX boxes of each type per CQ500 scan (initial manual labeling, bhx_*, and extrapolation to all series, bhxall_*), linked by StudyInstanceUID. '
                        'Derived from BHX (PhysioNet, CC BY-SA 4.0).'),
    ('MaskCheck', 'Model-free description of every released Seg-CQ500 mask: volume (|det(affine)|), attenuation inside the mask, overlap with the released brain mask and connected components.'),
    ('Adjudication', 'Stratified case list for independent neuroradiologist adjudication (coordinator copy with strata) and the randomized reading order (readers receive only the order).'),
    ('Registry', 'Values generated by the analysis code for the manuscript and Supplementary Information (key, value); reader labels of single scans are omitted. Intervals are 95% CIs as described in Methods.'),
    ('Checks', 'Formulas that recompute selected reported values from the sheets above, beside the value reported in the manuscript.'),
    ('License', 'CC BY 4.0, except for values derived from CQ500 images (the CQ500, SegCQ500, PublicModel, CQ500_UIDs, MaskCheck and Adjudication sheets and the CQ500 values in the Registry '
                'and Checks sheets), which keep the CQ500 license (CC BY-NC-SA 4.0), and the BHX_scan_counts sheet, derived from BHX (CC BY-SA 4.0). All values remain subject to the terms of their source datasets.'),
]

wb = Workbook()
ws = wb.active
ws.title = 'README'
ws['A1'] = 'Supplementary Data 1'
ws['A1'].font = Font(name=FONT, bold=True, size=13)
for i, (k, v) in enumerate(README[1:], start=3):
    ws.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
    c = ws.cell(row=i, column=2, value=v)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=i, column=1).alignment = Alignment(vertical='top')
ws.column_dimensions['A'].width = 18
ws.column_dimensions['B'].width = 120


def add_sheet(name, df, widths=None):
    sh = wb.create_sheet(name)
    for j, col in enumerate(df.columns, start=1):
        c = sh.cell(row=1, column=j, value=col)
        c.font = Font(name=FONT, bold=True, size=10)
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = Border(bottom=thin)
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            if hasattr(v, 'item'):
                v = v.item()
            if isinstance(v, bool):
                v = 'yes' if v else 'no'  # plain text, so that spreadsheet programs do not store booleans as formulas
            c = sh.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT, size=10)
    sh.freeze_panes = 'B2'
    for j, col in enumerate(df.columns, start=1):
        w = max(len(str(col)) * 0.9, 10)
        if df[col].dtype == object:
            w = max(w, min(40, df[col].astype(str).str.len().max() + 2))
        sh.column_dimensions[get_column_letter(j)].width = min(w, 42)
    sh.row_dimensions[1].height = 30
    return sh


add_sheet('CQ500', cq_out)
add_sheet('BHSD_test', bh)
add_sheet('QC_features', qc_out)
add_sheet('SegCQ500', seg)
add_sheet('MRBHSD', mr_out)
add_sheet('PHE_laterality', ph_out)
add_sheet('PHE_edema_test', ed_out)
pmd = pd.read_csv(E / 'public_model_cq500' / 'public_model_per_scan.csv')
add_sheet('PublicModel', pmd)
add_sheet('CQ500_UIDs', pd.read_csv(E / 'cq500_uids' / 'cq500_series_uids.csv'))
bxs = pd.read_csv(E / 'bhx_scan_labels' / 'bhx_scan_labels.csv')
add_sheet('BHX_scan_counts', bxs)
mkc = pd.read_csv(W / 'data' / 'derived' / 'segcq500_mask_check.csv')
add_sheet('MaskCheck', mkc)
adj = pd.read_csv(OUT / 'tables' / 'adjudication_case_list.csv').merge(pd.read_csv(OUT / 'tables' / 'adjudication_reading_order.csv'), on='scan')
add_sheet('Adjudication', adj[['reading_number', 'scan', 'stratum', 'description', 'pool_size']].sort_values('reading_number'))
reg_sh = add_sheet('Registry', reg_out)
reg_sh.column_dimensions['A'].width = 34
reg_sh.column_dimensions['B'].width = 90


def col(df, name):
    return get_column_letter(list(df.columns).index(name) + 1)


nb, ns, nc = len(bh) + 1, len(seg) + 1, len(cq_out) + 1
checks = [
    ('Mean Dice, v1, BHSD test (96 scans)', f"=AVERAGE(BHSD_test!{col(bh, 'v1_dice')}2:{col(bh, 'v1_dice')}{nb})", 0.436, '0.000'),
    ('Mean Dice, v2-PM, BHSD test', f"=AVERAGE(BHSD_test!{col(bh, 'v2PM_dice')}2:{col(bh, 'v2PM_dice')}{nb})", 0.390, '0.000'),
    ('Mean Dice, v2, BHSD test', f"=AVERAGE(BHSD_test!{col(bh, 'v2_dice')}2:{col(bh, 'v2_dice')}{nb})", 0.461, '0.000'),
    ('Mean v3b EDH Dice in scans with reference EDH', f"=AVERAGEIFS(BHSD_test!{col(bh, 'v3b_EDH_dice')}2:{col(bh, 'v3b_EDH_dice')}{nb},BHSD_test!{col(bh, 'EDH_reference_ml')}2:{col(bh, 'EDH_reference_ml')}{nb},\">0\")", 0.078, '0.000'),
    ('Mean v3b IPH Dice in scans with reference IPH', f"=AVERAGEIFS(BHSD_test!{col(bh, 'v3b_IPH_dice')}2:{col(bh, 'v3b_IPH_dice')}{nb},BHSD_test!{col(bh, 'IPH_reference_ml')}2:{col(bh, 'IPH_reference_ml')}{nb},\">0\")", 0.405, '0.000'),
    ('BHSD test scans flagged by the QC gate', f"=COUNTIF(BHSD_test!{col(bh, 'qc_gate_flag')}2:{col(bh, 'qc_gate_flag')}{nb},\"yes\")", 1, '0'),
    ('BHSD training scans flagged by the QC gate', f"=COUNTIFS(QC_features!B2:B189,\"train\",QC_features!F2:F189,\"yes\")", 0, '0'),
    ('Mean Dice, v2, Seg-CQ500 (43 scans)', f"=AVERAGE(SegCQ500!{col(seg, 'v2_dice')}2:{col(seg, 'v2_dice')}{ns})", 0.493, '0.000'),
    ('Mean Dice, BLAST-CT, Seg-CQ500', f"=AVERAGE(SegCQ500!{col(seg, 'BLASTCT_dice')}2:{col(seg, 'BLASTCT_dice')}{ns})", 0.531, '0.000'),
    ('Mean Dice, v1 against each rater, MR-BHSD (129 comparisons)', "=AVERAGE(MRBHSD!D2:D130)", 0.399, '0.000'),
    ('Mean edema Dice, 37 held-out patients', "=AVERAGE(PHE_edema_test!B2:B38)", 0.336, '0.000'),
    ('CQ500 scans called positive at 1 mL (true + false positives, 168 + 115)', f"=SUM(CQ500!{col(cq_out, 'v2_call_at_1ml')}2:{col(cq_out, 'v2_call_at_1ml')}{nc})", 283, '0'),
    ('CQ500 scans flagged by the QC gate', f"=COUNTIF(CQ500!{col(cq_out, 'qc_gate_flag')}2:{col(cq_out, 'qc_gate_flag')}{nc},\"yes\")", 0, '0'),
    ('CQ500 scans with an estimated ideal midline', f"=COUNTIF(CQ500!{col(cq_out, 'midline_status')}2:{col(cq_out, 'midline_status')}{nc},\"estimated\")", 487, '0'),
    ('CQ500 series with slice interval of 1.25 mm or less', f"=COUNTIF(CQ500!{col(cq_out, 'slice_interval_mm')}2:{col(cq_out, 'slice_interval_mm')}{nc},\"<=1.3\")", 25, '0'),
    ('Public model: scans with any-hemorrhage probability >= 0.5 (true + false positives)', f"=COUNTIF(PublicModel!{col(pmd, 'p_any')}2:{col(pmd, 'p_any')}{len(pmd) + 1},\">=0.5\")",
     PMJ['ops_all']['tp'] + PMJ['ops_all']['fp'], '0'),
    ('CQ500 scans with at least one BHX box (initial labeling)', f"=COUNTIF(BHX_scan_counts!{col(bxs, 'bhx_any')}2:{col(bxs, 'bhx_any')}{len(bxs) + 1},\">0\")", 226, '0'),
    ('Seg-CQ500 CT-108 mask: median attenuation, HU', f"=INDEX(MaskCheck!{col(mkc, 'hu_median')}2:{col(mkc, 'hu_median')}{len(mkc) + 1},MATCH(\"CQ500-CT-108\",MaskCheck!A2:A{len(mkc) + 1},0))", 23, '0'),
]
ck = wb.create_sheet('Checks')
for j, h in enumerate(['Check', 'Recomputed (formula)', 'Reported in manuscript', 'Difference'], start=1):
    c = ck.cell(row=1, column=j, value=h)
    c.font = Font(name=FONT, bold=True, size=10); c.fill = HFILL
for i, (lab, f, rep, fmt) in enumerate(checks, start=2):
    ck.cell(row=i, column=1, value=lab).font = Font(name=FONT, size=10)
    c = ck.cell(row=i, column=2, value=f); c.font = Font(name=FONT, size=10); c.number_format = fmt
    c = ck.cell(row=i, column=3, value=rep); c.font = Font(name=FONT, size=10, color='0000FF'); c.number_format = fmt
    c = ck.cell(row=i, column=4, value=f'=ROUND(B{i},3)-C{i}'); c.font = Font(name=FONT, size=10); c.number_format = '0.000'
ck.cell(row=len(checks) + 3, column=1, value='Blue values are the values reported in the manuscript; Difference = recomputed value rounded to three decimals minus the reported value.').font = Font(name=FONT, italic=True, size=9)
ck.column_dimensions['A'].width = 62
for L in 'BCD':
    ck.column_dimensions[L].width = 22
ck.freeze_panes = 'A2'

dst.parent.mkdir(parents=True, exist_ok=True)
wb.save(dst)
print('wrote', dst, {s: wb[s].max_row for s in wb.sheetnames})
